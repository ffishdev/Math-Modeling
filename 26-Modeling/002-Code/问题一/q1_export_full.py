"""把问题一的四项计算结果汇成可审阅的四工作表 Excel。

先运行 q1_solve.py 生成 q1_results.json，再运行本脚本。
仅用 py310 已有的 openpyxl；不覆盖竞赛提交模板或先前的 Q1 模板填写版。
"""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent


def find_project_root() -> Path:
    for parent in Path(__file__).resolve().parents:
        if (parent / "26-Modeling/001-Issue/D题").is_dir():
            return parent
    raise FileNotFoundError("无法从脚本位置向上找到 26-Modeling/001-Issue/D题")


ROOT = find_project_root()
RESULT = HERE / "q1_results.json"
SOURCE = ROOT / "26-Modeling/001-Issue/D题/数据/无人机应急物资运输基础数据/运输无人机数据.xlsx"
DEST = HERE / "问题一四项结果总览_Excel兼容版.xlsx"
PREVIOUS = HERE / "问题一四项结果总览.xlsx"

BLUE = "FF17365D"
PALE = "FFF3F7FA"
RED = "FFFCE8E6"
TEXT = "FF203040"


def model_specs():
    ws = load_workbook(SOURCE, read_only=True, data_only=True).active
    return {r[0]: {"mass": float(r[3]), "volume": float(r[4]), "energy": float(r[8])}
            for r in ws.iter_rows(values_only=True)
            if r[0] in {"A", "B", "C"} and isinstance(r[3], (int, float))}


def setup(ws, title, note, widths):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B6"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(widths))
    c = ws.cell(1, 1, title)
    c.font = Font(size=17, bold=True, color="FFFFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(widths))
    c = ws.cell(2, 1, note)
    c.font = Font(size=10, color=TEXT)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 32
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_options.horizontalCentered = True


def header(ws, row, labels):
    for col, value in enumerate(labels, 1):
        c = ws.cell(row, col, value)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.font = Font(size=10, bold=True, color="FFFFFFFF")
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def body(ws, row, values, formats=None, alert=False):
    for col, value in enumerate(values, 1):
        c = ws.cell(row, col, value)
        c.font = Font(size=10, color=TEXT)
        c.alignment = Alignment(vertical="center", wrap_text=(col == 4 and ws.title == "02_货箱组批"))
        c.fill = PatternFill("solid", fgColor=RED if alert else ("FFFFFFFF" if row % 2 else PALE))
        c.border = Border(bottom=Side(style="hair", color="FFDDE4E8"))
        if formats and col in formats and isinstance(value, (int, float)):
            c.number_format = formats[col]


def build_capacity(wb, data, specs):
    ws = wb.create_sheet("01_最大安全载荷")
    setup(ws, "第一项  三机型单点往返最大安全载荷",
          "返航安全余量按附件默认 20%；载荷取额定载重与能量可行上限的较小值。体积上限另在组批中检查。原始依据：节点、DEM、运输无人机参数。",
          [14, 18, 16, 18, 18, 18, 18])
    ws.cell(3, 1, "体积上限 m³")
    for col, model in zip((5, 6, 7), "ABC"):
        ws.cell(3, col, f"{model}: {specs[model]['volume']:.3f}")
    header(ws, 5, ["服务区", "水平距离 m", "航段最高 DEM m", "计划巡航海拔 m",
                   "A 最大安全载荷 kg", "B 最大安全载荷 kg", "C 最大安全载荷 kg"])
    lookup = {(r["reserve_pct"], r["zone"], r["model"]): r for r in data["capacities"]}
    for row, zone in enumerate(sorted(data["geometry"]), 6):
        g = data["geometry"][zone]
        values = [zone, g["distance_m"], g["max_dem_m"], g["cruise_altitude_m"]]
        values += [lookup[20.0, zone, model]["safe_mass_kg"] for model in "ABC"]
        body(ws, row, values, {2: "0.000", 3: "0.000", 4: "0.000", 5: "0.000", 6: "0.000", 7: "0.000"})
    ws.auto_filter.ref = "A5:G20"
    ws.print_title_rows = "1:5"


def build_batches(wb, data, specs):
    ws = wb.create_sheet("02_货箱组批")
    setup(ws, "第二项  80 箱单服务区货箱组批",
          "默认 20% 返航余量；各箱只安排一次，不跨服务区。往返时间为飞行时间；作业时间另含准备、装载和交接。此表与竞赛提交模板 Q1 工作表对应。",
          [13, 12, 10, 72, 11, 13, 14, 16, 16, 16, 16, 17, 18, 18])
    labels = ["架次编号", "服务区", "机型", "货箱编号列表", "箱数", "总质量 kg", "总体积 m³",
              "往返飞行 s", "累计作业 s", "架次能耗 kWh", "返航 SOC %",
              "额定载重 kg", "装载体积上限 m³", "电量余量百分点"]
    header(ws, 5, labels)
    trips = data["scenarios"]["20pct_lexicographic"]["trips"]
    for row, t in enumerate(trips, 6):
        m = specs[t["model"]]
        values = [t["trip_id"], t["zone"], t["model"], ",".join(t["box_ids"]), len(t["box_ids"]),
                  t["mass_kg"], t["volume_m3"], t["flight_s"], t["operation_s"],
                  t["energy_kwh"], t["return_soc_pct"], m["mass"], m["volume"], t["return_soc_pct"] - 20]
        body(ws, row, values, {6: "0.000", 7: "0.000", 8: "0.000", 9: "0.000", 10: "0.000000",
                               11: "0.0000", 12: "0.000", 13: "0.000", 14: "0.0000"})
        ws.row_dimensions[row].height = 45 if len(t["box_ids"]) > 5 else 30
    summary = data["scenarios"]["20pct_lexicographic"]["metrics"]
    last = len(trips) + 7
    body(ws, last, ["合计", "15 个服务区", "", "80 箱", 80, None, None, summary["flight_s"],
                    summary["operation_s"], summary["energy_kwh"]],
         {8: "0.000", 9: "0.000", 10: "0.000000"})
    for c in ws[last]:
        c.font = Font(size=10, bold=True, color=BLUE)
    ws.auto_filter.ref = f"A5:N{len(trips)+5}"
    ws.print_title_rows = "1:5"


def build_tradeoff(wb, data):
    ws = wb.create_sheet("03_目标权衡")
    setup(ws, "第三项  架次 能耗 作业时间权衡",
          "完成全部货箱交付后比较三个独立优化口径。推荐口径：先最少架次，再最低能耗，最后最少累计作业时间；无实体机调度，因此作业时间不是完工时刻。",
          [25, 15, 20, 19, 19, 24])
    header(ws, 5, ["20% 余量优化口径", "往返架次", "总能耗 kWh", "累计飞行 s", "累计作业 s", "与推荐方案的区别"])
    name = {"20pct_lexicographic": "架次优先 推荐", "20pct_energy": "能耗优先", "20pct_time": "作业时间优先"}
    base = data["scenarios"]["20pct_lexicographic"]["metrics"]
    for row, key in enumerate(name, 6):
        m = data["scenarios"][key]["metrics"]
        difference = "基准" if key == "20pct_lexicographic" else (
            "S008 分两批；多 1 架次" if key == "20pct_energy" else "与推荐方案相同")
        body(ws, row, [name[key], m["trip_count"], m["energy_kwh"], m["flight_s"], m["operation_s"], difference],
             {3: "0.000000", 4: "0.000", 5: "0.000"})
    ws.cell(11, 1, "能耗优先相对推荐方案")
    other = data["scenarios"]["20pct_energy"]["metrics"]
    ws.cell(12, 1, "节省能耗 kWh")
    ws.cell(12, 2, base["energy_kwh"] - other["energy_kwh"]).number_format = "0.000000"
    ws.cell(13, 1, "增加架次")
    ws.cell(13, 2, other["trip_count"] - base["trip_count"])
    ws.cell(14, 1, "增加累计作业 s")
    ws.cell(14, 2, other["operation_s"] - base["operation_s"]).number_format = "0.000"
    ws.merge_cells("A17:F17")
    ws["A17"] = "18 架次的下界：15 个服务区各至少 1 架次，S001、S002、S003 均至少再加 1 架次。"
    ws["A17"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[17].height = 30


def build_sensitivity(wb, data):
    ws = wb.create_sheet("04_余量敏感性")
    setup(ws, "第四项  返航安全余量敏感性",
          "安全载荷单位 kg；不可达表示该机型空载直接往返也达不到所设余量。40% 时全量交付不可行，不能把空白理解为零架次。",
          [13] + [15] * 12)
    header(ws, 5, ["返航余量", "最少架次", "总能耗 kWh", "累计作业 s", "全量交付", "变化说明"])
    for row, reserve in enumerate((10, 20, 30, 40), 6):
        key = f"{reserve}pct_lexicographic"
        s = data["scenarios"][key]
        m = s.get("metrics")
        note = {10: "与 20% 同组批", 20: "附件默认值", 30: "S004 和 S008 各增加 1 架次",
                40: "S004 饮用水单箱不可送；S008 空载也不可达"}[reserve]
        body(ws, row, [f"{reserve}%", None if m is None else m["trip_count"],
                       None if m is None else m["energy_kwh"],
                       None if m is None else m["operation_s"],
                       "不可行" if m is None else "可行", note],
             {3: "0.000000", 4: "0.000"}, alert=m is None)
    ws.cell(11, 1, "各机型在各服务区的最大安全载荷 kg")
    ws["A11"].font = Font(size=12, bold=True, color=BLUE)
    labels = ["服务区"] + [f"{reserve}% {model}" for reserve in (10, 20, 30, 40) for model in "ABC"]
    header(ws, 12, labels)
    lookup = {(r["reserve_pct"], r["zone"], r["model"]): r["safe_mass_kg"] for r in data["capacities"]}
    zones = sorted(data["geometry"])
    for row, zone in enumerate(zones, 13):
        values = [zone] + [lookup[float(reserve), zone, model] for reserve in (10, 20, 30, 40) for model in "ABC"]
        values = ["不可达" if value is None else value for value in values]
        body(ws, row, values, {col: "0.000" for col in range(2, 14)}, alert=False)
        for col in range(2, 14):
            if values[col-1] == "不可达":
                ws.cell(row, col).fill = PatternFill("solid", fgColor=RED)
    ws.auto_filter.ref = "A12:M27"
    ws.cell(30, 1, "各服务区组批架次数")
    ws["A30"].font = Font(size=12, bold=True, color=BLUE)
    header(ws, 31, ["服务区", "10%", "20%", "30%"])
    counts = {reserve: Counter(t["zone"] for t in data["scenarios"][f"{reserve}pct_lexicographic"]["trips"])
              for reserve in (10, 20, 30)}
    for row, zone in enumerate(zones, 32):
        body(ws, row, [zone] + [counts[reserve][zone] for reserve in (10, 20, 30)])
    ws.print_title_rows = "1:12"


def verify(path, data):
    wb = load_workbook(path, read_only=True, data_only=True)
    assert wb.sheetnames == ["01_最大安全载荷", "02_货箱组批", "03_目标权衡", "04_余量敏感性"]
    cap = wb.worksheets[0]
    assert cap["A6"].value == "S001" and cap["G20"].value == 80
    batches = wb.worksheets[1]
    trips = data["scenarios"]["20pct_lexicographic"]["trips"]
    assert len(trips) == 18
    box_ids = []
    for row_num, t in enumerate(trips, 6):
        assert batches.cell(row_num, 1).value == t["trip_id"]
        assert batches.cell(row_num, 4).value.split(",") == t["box_ids"]
        assert abs(batches.cell(row_num, 10).value - t["energy_kwh"]) < 1e-10
        box_ids.extend(t["box_ids"])
    assert len(box_ids) == len(set(box_ids)) == 80
    assert wb.worksheets[2]["B6"].value == 18
    sensitivity = wb.worksheets[3]
    assert sensitivity["E9"].value == "不可行"
    for i, zone in enumerate(sorted(data["geometry"]), 13):
        assert sensitivity.cell(i, 1).value == zone
    if PREVIOUS.exists():
        previous = load_workbook(PREVIOUS, read_only=True, data_only=True)
        assert previous.sheetnames == wb.sheetnames
        for new_sheet, old_sheet in zip(wb.worksheets, previous.worksheets):
            assert list(new_sheet.values) == list(old_sheet.values), new_sheet.title
    styled = load_workbook(path, read_only=True, data_only=False)
    for sheet in styled:
        for row in sheet:
            for cell in row:
                font = cell.font
                if font is None:
                    continue
                assert font.name != "PingFang SC"
                if font.color and font.color.type == "rgb":
                    assert font.color.rgb.startswith("FF")
    print(f"已核对四个工作表、18 架次、80 个唯一货箱及敏感性表：{path}")


def main():
    data = json.loads(RESULT.read_text(encoding="utf-8"))
    specs = model_specs()
    wb = Workbook()
    wb.remove(wb.active)
    build_capacity(wb, data, specs)
    build_batches(wb, data, specs)
    build_tradeoff(wb, data)
    build_sensitivity(wb, data)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(DEST)
    verify(DEST, data)


if __name__ == "__main__":
    main()
